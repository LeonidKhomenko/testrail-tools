#!/usr/bin/env python3
"""
TestRail Case History Analyzer - Version 5.4.7
(2-stage scan: non-CI first, CI only for unresolved automated cases + checkpoints + in-memory tests cache)

Key behavior:
- Full-history scan (no truncation by latest Nightly plans)
- Stage 1: scan NON-CI runs for all candidate cases
- Stage 2: scan CI runs only for cases with automation_id that still have no history after Stage 1
- Optimizations:
  - get_tests status shortcut (status != Untested => has history)
  - get_results_for_case only for candidate+Untested pairs
  - skip already-confirmed cases
  - early-stop only when all candidate cases (for this stage) have history
  - in-memory cache for get_tests pages (optional, enabled by default)
- Reliability:
  - checkpoint CSV export every N completed runs (default N=50, enabled by default)
- Automation ID detection:
  - uses confirmed key: custom_case_automation_id
  - raises explicit error if that key is missing in any case payload
"""

import csv
import re
import sys
import time
import threading
from base64 import b64encode
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from email.utils import parsedate_to_datetime
from getpass import getpass
from urllib.parse import urlsplit

import requests

UNTESTED_STATUS_ID = 3
AUTOMATION_ID_FIELD = 'custom_case_automation_id'


def extract_automation_id(case_obj):
    if AUTOMATION_ID_FIELD not in case_obj:
        raise ValueError(
            f'Case {case_obj.get("id")}: API field {AUTOMATION_ID_FIELD!r} is missing. '
            'Cannot safely decide whether to skip CI scanning.'
        )

    value = case_obj[AUTOMATION_ID_FIELD]
    if value is None:
        return ''
    if not isinstance(value, str):
        raise ValueError(
            f'Case {case_obj.get("id")}: unexpected Automation ID type: {type(value).__name__}'
        )
    return value.strip()


class TestRailAPI:
    def __init__(self, base_url, email, api_token):
        self.base_url = base_url.strip().rstrip('/')
        parsed = urlsplit(self.base_url)
        if parsed.scheme not in ('http', 'https') or not parsed.hostname:
            raise ValueError('Invalid TestRail URL. Example: https://renesas.testrail.io')

        auth = b64encode(f'{email}:{api_token}'.encode()).decode()
        self._headers = {
            'Authorization': f'Basic {auth}',
            'Content-Type': 'application/json',
        }

        self._local = threading.local()
        self._rate_lock = threading.Lock()
        self._request_interval = 0.4
        self._next_request_at = 0.0
        self._blocked_until = 0.0

    def _get_session(self):
        if not hasattr(self._local, 'session'):
            s = requests.Session()
            s.headers.update(self._headers)
            self._local.session = s
        return self._local.session

    def _wait_for_request_slot(self):
        while True:
            with self._rate_lock:
                now = time.monotonic()
                allowed_at = max(self._next_request_at, self._blocked_until)
                delay = allowed_at - now
                if delay <= 0:
                    self._next_request_at = now + self._request_interval
                    return
            time.sleep(delay)

    def _retry_after_seconds(self, response, fallback):
        value = response.headers.get('Retry-After')
        if value:
            try:
                return max(0.0, float(value))
            except ValueError:
                try:
                    dt = parsedate_to_datetime(value)
                    if dt.tzinfo is None:
                        dt = dt.replace(tzinfo=timezone.utc)
                    return max(0.0, dt.timestamp() - time.time())
                except Exception:
                    pass

        m = re.search(r'retry after\s+(\d+(?:\.\d+)?)\s*seconds?', response.text or '', re.IGNORECASE)
        if m:
            return float(m.group(1))

        return fallback

    def api_request(self, method, endpoint, json_data=None):
        method = method.upper()
        if method not in ('GET', 'POST'):
            raise ValueError(f'Unsupported method: {method}')

        url = f'{self.base_url}/index.php?/api/v2{endpoint}'
        session = self._get_session()
        max_attempts = 6

        for attempt in range(max_attempts):
            self._wait_for_request_slot()

            try:
                if method == 'GET':
                    response = session.get(url, timeout=60)
                else:
                    response = session.post(url, json=json_data, timeout=60)
            except (requests.exceptions.Timeout, requests.exceptions.ConnectionError):
                if method != 'GET' or attempt == max_attempts - 1:
                    raise
                delay = min(2 ** (attempt + 1), 30)
                print(f'⚠️ Network error; retrying GET in {delay}s ({attempt + 1}/{max_attempts - 1}) for {endpoint}')
                time.sleep(delay)
                continue

            retryable = method == 'GET' and (
                response.status_code == 429 or response.status_code in (500, 502, 503, 504)
            )
            if retryable and attempt < max_attempts - 1:
                fallback = 60 if response.status_code == 429 else min(2 ** (attempt + 1), 30)
                delay = self._retry_after_seconds(response, fallback) + 1.0

                with self._rate_lock:
                    self._blocked_until = max(self._blocked_until, time.monotonic() + delay)

                print(f'⏳ HTTP {response.status_code}: pausing API requests for {delay:.1f}s; retry {attempt + 1}/{max_attempts - 1}')
                response.close()
                continue

            try:
                response.raise_for_status()
                return response.json() if response.text else None
            finally:
                response.close()

    def get_projects(self):
        return self.api_request('GET', '/get_projects')

    def get_suites(self, project_id):
        return self.api_request('GET', f'/get_suites/{project_id}')

    def get_sections(self, project_id, suite_id=None, offset=0, limit=250):
        endpoint = f'/get_sections/{project_id}&offset={offset}&limit={limit}'
        if suite_id is not None:
            endpoint += f'&suite_id={suite_id}'
        return self.api_request('GET', endpoint)

    def get_cases(self, project_id, suite_id=None, section_id=None, offset=0, limit=250):
        endpoint = f'/get_cases/{project_id}&offset={offset}&limit={limit}'
        if suite_id is not None:
            endpoint += f'&suite_id={suite_id}'
        if section_id is not None:
            endpoint += f'&section_id={section_id}'
        return self.api_request('GET', endpoint)

    def get_plans(self, project_id, offset=0, limit=250):
        return self.api_request('GET', f'/get_plans/{project_id}&offset={offset}&limit={limit}')

    def get_plan(self, plan_id):
        return self.api_request('GET', f'/get_plan/{plan_id}')

    def get_runs(self, project_id, suite_id=None, is_completed=None, offset=0, limit=250):
        endpoint = f'/get_runs/{project_id}&offset={offset}&limit={limit}'
        if suite_id is not None:
            endpoint += f'&suite_id={suite_id}'
        if is_completed is not None:
            endpoint += f'&is_completed={1 if is_completed else 0}'
        return self.api_request('GET', endpoint)

    def get_tests(self, run_id, offset=0, limit=250):
        return self.api_request('GET', f'/get_tests/{run_id}&offset={offset}&limit={limit}')

    def get_results_for_case(self, run_id, case_id, offset=0, limit=250):
        return self.api_request('GET', f'/get_results_for_case/{run_id}/{case_id}&offset={offset}&limit={limit}')


def normalize_list_response(response, key):
    if isinstance(response, list):
        return response
    if isinstance(response, dict):
        val = response.get(key, [])
        return val if isinstance(val, list) else []
    return []


def require_list(response, key):
    if isinstance(response, list):
        items = response
    elif isinstance(response, dict) and key in response:
        items = response[key]
    else:
        raise ValueError(f'Unexpected API response: expected list or "{key}" field')

    if not isinstance(items, list):
        raise ValueError(f'API field "{key}" is not a list')

    if any(not isinstance(item, dict) for item in items):
        raise ValueError(f'API field "{key}" contains invalid items')

    return items


def fetch_all_sections(client, project_id, suite_id=None):
    sections = []
    offset = 0
    limit = 250
    while True:
        response = client.get_sections(project_id, suite_id=suite_id, offset=offset, limit=limit)
        chunk = normalize_list_response(response, 'sections')
        if not chunk:
            break
        sections.extend(chunk)
        if len(chunk) < limit:
            break
        offset += limit
    return sections


def build_section_maps(sections):
    by_id = {s['id']: s for s in sections}
    children = {}
    for s in sections:
        pid = s.get('parent_id')
        children.setdefault(pid, []).append(s['id'])

    path_cache = {}

    def build_path(sid):
        if sid in path_cache:
            return path_cache[sid]
        sec = by_id[sid]
        pid = sec.get('parent_id')
        if pid and pid in by_id:
            path = f'{build_path(pid)} > {sec["name"]}'
        else:
            path = sec['name']
        path_cache[sid] = path
        return path

    section_path = {sid: build_path(sid) for sid in by_id}
    return by_id, children, section_path


def collect_descendants(children_map, roots):
    result = set()
    stack = list(roots)
    while stack:
        sid = stack.pop()
        if sid in result:
            continue
        result.add(sid)
        stack.extend(children_map.get(sid, []))
    return result


def parse_section_selection(user_input):
    raw = user_input.replace(',', ' ').split()
    return [int(token) for token in raw]


def choose_sections_interactive(by_id, children_map, section_path):
    selected = set()
    current_parent = None

    def current_children(pid):
        ids = children_map.get(pid, [])
        return sorted(ids, key=lambda sid: by_id[sid]['name'].lower())

    while True:
        kids = current_children(current_parent)
        if current_parent is None:
            print('\n=== TOP-LEVEL SECTIONS ===')
        else:
            print(f'\n=== SUBSECTIONS OF: {section_path[current_parent]} ===')

        if not kids:
            print('  (No child sections)')
        else:
            for idx, sid in enumerate(kids, 1):
                mark = '✓' if sid in selected else ' '
                print(f'  {idx:>2}. [{mark}] {by_id[sid]["name"]} (ID: {sid})')

        print('\nActions:')
        print('  a  - add ALL sections from current level (with descendants)')
        print('  d  - drill down into one subsection')
        print('  s  - select one/many sections from current level (with descendants)')
        print('  c  - select CURRENT section fully (current + descendants)')
        print('  u  - go UP one level')
        print('  l  - list selected sections')
        print('  f  - finish selection')
        print('  x  - use ALL sections (no filtering)')

        action = input('\nChoose action [a/d/s/c/u/l/f/x]: ').strip().lower()

        if action == 'x':
            return None
        if action == 'f':
            if selected:
                return selected
            print('⚠️ Nothing selected yet.')
            continue
        if action == 'l':
            if not selected:
                print('  (No sections selected)')
            else:
                print('\nSelected sections:')
                for sid in sorted(selected):
                    print(f'  - {sid}: {section_path[sid]}')
            continue
        if action == 'u':
            if current_parent is None:
                print('⚠️ Already at top level.')
                continue
            parent = by_id[current_parent].get('parent_id')
            current_parent = parent if parent in by_id else None
            continue
        if action == 'c':
            if current_parent is None:
                print('⚠️ At root there is no single current section. Use "a" or "s".')
                continue
            desc = collect_descendants(children_map, [current_parent])
            selected.update(desc)
            print(f'✅ Added branch: {section_path[current_parent]} ({len(desc)} sections)')
            continue
        if action == 'a':
            if not kids:
                print('⚠️ No sections on this level.')
                continue
            added = 0
            for sid in kids:
                desc = collect_descendants(children_map, [sid])
                before = len(selected)
                selected.update(desc)
                added += len(selected) - before
            print(f'✅ Added all current-level branches (+{added})')
            continue
        if action == 'd':
            if not kids:
                print('⚠️ No subsection to drill into.')
                continue
            raw = input('Enter subsection number: ').strip()
            try:
                idx = int(raw)
                if idx < 1 or idx > len(kids):
                    raise ValueError
                current_parent = kids[idx - 1]
            except ValueError:
                print('❌ Invalid number.')
            continue
        if action == 's':
            if not kids:
                print('⚠️ No sections on this level.')
                continue
            raw = input('Enter numbers (comma/space), e.g. 1,3,5: ').strip()
            try:
                nums = parse_section_selection(raw)
                bad = [n for n in nums if n < 1 or n > len(kids)]
                if bad:
                    print(f'❌ Invalid numbers: {bad}')
                    continue
                added = 0
                for n in nums:
                    sid = kids[n - 1]
                    desc = collect_descendants(children_map, [sid])
                    before = len(selected)
                    selected.update(desc)
                    added += len(selected) - before
                print(f'✅ Added selected branches (+{added})')
            except ValueError:
                print('❌ Invalid input format.')
            continue

        print('❌ Unknown action.')


def fetch_all_cases(client, project_id, suite_id=None, selected_section_ids=None):
    if selected_section_ids is None:
        section_ids = [None]
    else:
        section_ids = sorted(set(selected_section_ids))
        if not section_ids:
            return []

    cases_by_id = {}
    started = time.monotonic()

    for i, section_id in enumerate(section_ids, 1):
        label = 'Entire suite' if section_id is None else f'Section {section_id} ({i}/{len(section_ids)})'
        offset = 0
        limit = 250
        seen_ids = set()
        section_count = 0

        while True:
            print(f'  ⏳ {label}: requesting offset={offset}...', flush=True)
            response = client.get_cases(project_id, suite_id=suite_id, section_id=section_id, offset=offset, limit=limit)

            if isinstance(response, list):
                chunk = response
            elif isinstance(response, dict):
                chunk = response.get('cases')
            else:
                chunk = None

            if not isinstance(chunk, list) or any(not isinstance(c, dict) for c in chunk):
                raise ValueError(f'{label}: invalid cases response; stopped')

            page_ids = set()
            for case in chunk:
                cid = case['id']
                if cid in page_ids or cid in seen_ids:
                    raise ValueError(f'{label}: duplicate case ID {cid} during pagination')
                if section_id is not None and case.get('section_id') != section_id:
                    raise ValueError(f'{label}: case {cid} returned from different section')
                page_ids.add(cid)
                cases_by_id[cid] = case

            seen_ids.update(page_ids)
            section_count += len(chunk)

            print(
                f'  ✓ {label}: {section_count} cases; total unique: {len(cases_by_id)}; '
                f'elapsed: {time.monotonic() - started:.1f}s',
                flush=True
            )

            if isinstance(response, dict) and '_links' in response:
                links = response['_links']
                if not isinstance(links, dict) or 'next' not in links:
                    raise ValueError(f'{label}: invalid pagination metadata')
                has_next = bool(links['next'])
            else:
                has_next = len(chunk) == limit

            if not has_next:
                break
            if not chunk:
                raise ValueError(f'{label}: empty page with next link present')
            offset += len(chunk)

    return sorted(cases_by_id.values(), key=lambda c: c['id'])


def fetch_runs_from_plans(client, project_id, suite_id=None):
    runs = {}
    run_to_plan = {}
    offset = 0
    limit = 250

    while True:
        plans_resp = client.get_plans(project_id, offset=offset, limit=limit)
        plans = require_list(plans_resp, 'plans')
        if not plans:
            break

        for p in plans:
            plan_id = p['id']
            plan_name = p.get('name', f'Plan {plan_id}')

            detail = client.get_plan(plan_id)
            if not isinstance(detail, dict):
                raise ValueError(f'Plan {plan_id}: invalid plan details response')
            entries = detail.get('entries')
            if not isinstance(entries, list):
                raise ValueError(f'Plan {plan_id}: missing/invalid entries')

            for entry in entries:
                if not isinstance(entry, dict):
                    raise ValueError(f'Plan {plan_id}: invalid entry item')
                entry_suite_id = entry.get('suite_id')
                if suite_id is not None and entry_suite_id != suite_id:
                    continue

                entry_runs = entry.get('runs')
                if not isinstance(entry_runs, list):
                    raise ValueError(f'Plan {plan_id}: invalid runs in entry')

                for run in entry_runs:
                    rid = run.get('id')
                    if not isinstance(rid, int):
                        raise ValueError(f'Plan {plan_id}: run without valid id')
                    if rid not in runs:
                        runs[rid] = {
                            'id': rid,
                            'name': run.get('name', f'Run {rid}'),
                            'suite_id': entry_suite_id,
                            'is_completed': run.get('is_completed', False),
                            'created_on': run.get('created_on'),
                            'source': f'plan:{plan_name}',
                        }

                    run_to_plan[rid] = {
                        'plan_id': plan_id,
                        'plan_name': plan_name,
                        'plan_created_on': p.get('created_on'),
                    }

        if len(plans) < limit:
            break
        offset += limit

    return runs, run_to_plan


def fetch_standalone_runs(client, project_id, suite_id=None):
    runs = {}
    offset = 0
    limit = 250

    while True:
        resp = client.get_runs(project_id, suite_id=suite_id, is_completed=None, offset=offset, limit=limit)
        chunk = require_list(resp, 'runs')
        if not chunk:
            break

        for run in chunk:
            rid = run.get('id')
            if not isinstance(rid, int):
                raise ValueError('Standalone run without valid id')
            runs[rid] = {
                'id': rid,
                'name': run.get('name', f'Run {rid}'),
                'suite_id': run.get('suite_id'),
                'is_completed': run.get('is_completed', False),
                'created_on': run.get('created_on'),
                'source': 'standalone',
            }

        if len(chunk) < limit:
            break
        offset += limit

    return runs


def is_ci_plan_name(plan_name):
    name = str(plan_name or '').strip()
    if re.match(r'^\s*Nightly\s+CI\s*-\s*\d{4}-\d{2}-\d{2}\s*$', name, re.IGNORECASE):
        return True
    return bool(re.search(r'\bnightly\s*ci\b', name, re.IGNORECASE))


def split_runs_ci_non_ci(all_runs, run_to_plan):
    non_ci_runs = []
    ci_runs = []

    for r in all_runs:
        rid = r['id']
        pm = run_to_plan.get(rid)
        if pm is None:
            non_ci_runs.append(r)
            continue

        pname = pm.get('plan_name')
        if is_ci_plan_name(pname):
            ci_runs.append(r)
        else:
            non_ci_runs.append(r)

    print('\n📌 Run groups:')
    print(f'   Non-CI runs: {len(non_ci_runs)}')
    print(f'   CI runs: {len(ci_runs)}')
    print(f'   Total: {len(non_ci_runs) + len(ci_runs)}')
    return non_ci_runs, ci_runs


def save_checkpoint_csv(base_name, stage_label, done_runs, candidate_cases, history_ids, in_open_ids):
    rows = build_final_report(candidate_cases, history_ids, in_open_ids)
    filename = f'{base_name}_checkpoint_{stage_label}_run{done_runs}.csv'
    save_to_csv(rows, filename)
    print(f'💾 Checkpoint saved: {filename}')


def scan_runs_newest_first(
    client,
    candidate_cases,
    all_runs,
    workers=4,
    history_ids=None,
    in_open_ids=None,
    stage_label='stage',
    checkpoint_enabled=True,
    checkpoint_every_runs=50,
    checkpoint_base_name='testrail_checkpoint',
    tests_cache_enabled=True,
    tests_cache=None
):
    candidate_ids = {c['id'] for c in candidate_cases}
    history_ids = history_ids if history_ids is not None else set()
    in_open_ids = in_open_ids if in_open_ids is not None else set()
    lock = threading.Lock()
    stop = threading.Event()

    tests_cache = tests_cache if tests_cache is not None else {}
    tests_cache_lock = threading.Lock()

    def recency(run):
        c = run.get('created_on')
        if isinstance(c, (int, float)):
            return int(c)
        return int(run.get('id', 0))

    runs = sorted(all_runs, key=lambda r: -recency(r))

    def get_tests_cached(run_id, offset, limit):
        key = (run_id, offset, limit)
        if tests_cache_enabled:
            with tests_cache_lock:
                if key in tests_cache:
                    return tests_cache[key]
        resp = client.get_tests(run_id, offset=offset, limit=limit)
        if tests_cache_enabled:
            with tests_cache_lock:
                tests_cache[key] = resp
        return resp

    def scan_run(run):
        rid = run['id']
        is_open = run.get('is_completed') is not True
        offset, limit = 0, 250
        seen_test_ids = set()
        untested_ids = set()

        while not stop.is_set():
            with lock:
                if all(cid in history_ids for cid in candidate_ids):
                    stop.set()
                    return

            resp = get_tests_cached(rid, offset=offset, limit=limit)
            tests = require_list(resp, 'tests')

            page_ids = {t.get('id') for t in tests}
            if None in page_ids or len(page_ids) != len(tests) or (page_ids & seen_test_ids):
                raise ValueError(f'Run {rid}: duplicate/malformed test pagination')
            seen_test_ids.update(page_ids)

            for t in tests:
                cid = t.get('case_id')
                if cid not in candidate_ids:
                    continue

                with lock:
                    if is_open:
                        in_open_ids.add(cid)
                    if cid in history_ids:
                        continue

                    status_id = t.get('status_id')
                    if status_id not in (None, UNTESTED_STATUS_ID):
                        history_ids.add(cid)
                        continue

                untested_ids.add(cid)

            if stop.is_set():
                return

            if isinstance(resp, dict) and '_links' in resp:
                links = resp['_links']
                if not isinstance(links, dict) or 'next' not in links:
                    raise ValueError(f'Run {rid}: invalid _links pagination metadata')
                has_next = bool(links['next'])
            else:
                has_next = len(tests) == limit

            if not has_next:
                break
            if not tests:
                raise ValueError(f'Run {rid}: empty page with next link')
            offset += len(tests)

        for cid in sorted(untested_ids):
            if stop.is_set():
                return

            with lock:
                if cid in history_ids:
                    continue

            resp = client.get_results_for_case(rid, cid, offset=0, limit=1)
            results = require_list(resp, 'results')
            if results:
                with lock:
                    history_ids.add(cid)

        with lock:
            if all(cid in history_ids for cid in candidate_ids):
                stop.set()

    total = len(runs)
    print(f'🔎 Scanning {total} runs (newest first) for {len(candidate_ids)} cases using {workers} workers...')
    if tests_cache_enabled:
        print('🧠 get_tests cache: enabled')
    else:
        print('🧠 get_tests cache: disabled')

    futures = {}
    executor = ThreadPoolExecutor(max_workers=workers)

    try:
        futures = {executor.submit(scan_run, r): r['id'] for r in runs}
        for done, fut in enumerate(as_completed(futures), 1):
            rid = futures[fut]
            try:
                fut.result()
            except Exception as exc:
                raise RuntimeError(
                    f'Run {rid}: scan failed ({exc}). No final report will be generated.'
                ) from exc

            with lock:
                hist_count = len(history_ids)
                all_found = all(cid in history_ids for cid in candidate_ids)

            if done % 25 == 0 or done == total or all_found:
                print(f'  ✓ Runs done: {done}/{total}; shared cases with history: {hist_count}')

            if checkpoint_enabled and checkpoint_every_runs > 0 and (done % checkpoint_every_runs == 0):
                try:
                    save_checkpoint_csv(
                        base_name=checkpoint_base_name,
                        stage_label=stage_label,
                        done_runs=done,
                        candidate_cases=candidate_cases,
                        history_ids=history_ids,
                        in_open_ids=in_open_ids
                    )
                except Exception as cp_exc:
                    print(f'⚠️ Checkpoint save failed at run {done}: {cp_exc}')

            if all_found:
                print('  ✅ Early stop for this stage: history found for all stage candidate cases.')
                stop.set()
                break
    finally:
        stop.set()
        for f in futures:
            f.cancel()
        executor.shutdown(wait=True)

    return history_ids, in_open_ids


def build_final_report(candidate_cases, history_ids, in_open_ids):
    rows = []
    for c in candidate_cases:
        cid = c['id']
        if cid in history_ids:
            status = 'has_history'
        elif cid in in_open_ids:
            status = 'in_open_run_no_results'
        else:
            status = 'never_executed'

        rows.append({
            'case_id': cid,
            'link': c['link'],
            'title': c['title'],
            'section': c['section'],
            'automation_id': c.get('automation_id', ''),
            'has_automation_id': 'YES' if c.get('has_automation_id') else 'NO',
            'status': status,
            'safe_to_delete_candidate': 'YES' if status == 'never_executed' else 'NO'
        })

    rows.sort(key=lambda r: r['case_id'])
    return rows


def escape_spreadsheet_formula(value):
    if not isinstance(value, str):
        return value
    stripped = value.lstrip()
    if stripped.startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def save_to_csv(rows, filename):
    if not rows:
        print(f'⚠️ No data to save for {filename}')
        return

    fieldnames = ['case_id', 'link', 'title', 'section', 'automation_id', 'has_automation_id', 'status', 'safe_to_delete_candidate']
    with open(filename, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        for row in rows:
            safe_row = dict(row)
            safe_row['title'] = escape_spreadsheet_formula(safe_row.get('title', ''))
            safe_row['section'] = escape_spreadsheet_formula(safe_row.get('section', ''))
            safe_row['automation_id'] = escape_spreadsheet_formula(str(safe_row.get('automation_id', '')))
            w.writerow(safe_row)

    print(f'✅ CSV saved: {filename}')


def save_to_xlsx(rows, filename):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        print('⚠️ openpyxl not installed, fallback to CSV.')
        save_to_csv(rows, filename.replace('.xlsx', '.csv'))
        return

    if not rows:
        print('⚠️ No data to save')
        return

    wb = Workbook()
    ws = wb.active
    ws.title = 'Case History Analysis'

    headers = ['Case ID', 'Link', 'Title', 'Section', 'Automation ID', 'Has Automation ID', 'Status', 'Safe To Delete Candidate']
    ws.append(headers)

    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    for c in ws[1]:
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal='center', vertical='center')

    for r in rows:
        ws.append([
            r['case_id'],
            r['link'],
            escape_spreadsheet_formula(r['title']),
            escape_spreadsheet_formula(r['section']),
            escape_spreadsheet_formula(str(r.get('automation_id', ''))),
            r.get('has_automation_id', 'NO'),
            r['status'],
            r['safe_to_delete_candidate']
        ])

    for row_idx in range(2, len(rows) + 2):
        cell = ws[f'B{row_idx}']
        cell.hyperlink = cell.value
        cell.style = 'Hyperlink'

    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 65
    ws.column_dimensions['C'].width = 60
    ws.column_dimensions['D'].width = 60
    ws.column_dimensions['E'].width = 28
    ws.column_dimensions['F'].width = 20
    ws.column_dimensions['G'].width = 32
    ws.column_dimensions['H'].width = 28

    wb.save(filename)
    print(f'✅ XLSX saved: {filename}')


def safe_filename_base(name, fallback):
    cleaned = ''.join(ch for ch in name if ch not in '<>:"/\\|?*').strip().strip('.')
    return cleaned or fallback


def main():
    print('🚀 TestRail Case History Analyzer (Version 5.4.7: checkpoints + tests cache)\n')

    default_url = 'https://renesas.testrail.io'
    default_email = 'leonid.khomenko.bx@bp.renesas.com'

    testrail_url = input(f'Enter TestRail URL [{default_url}]: ').strip() or default_url
    email = input(f'Enter TestRail email [{default_email}]: ').strip() or default_email
    api_token = getpass('Enter TestRail API token: ').strip()

    try:
        client = TestRailAPI(testrail_url, email, api_token)
    except ValueError as e:
        print(f'❌ {e}')
        sys.exit(1)

    print('\n📦 Fetching projects...')
    projects = normalize_list_response(client.get_projects(), 'projects')
    if not projects:
        print('❌ No projects found')
        sys.exit(1)

    print('\nAvailable projects:')
    for i, p in enumerate(projects, 1):
        print(f'  {i}. {p["name"]} (ID: {p["id"]})')

    try:
        p_idx = int(input('\nSelect project number: ').strip()) - 1
        project = projects[p_idx]
    except (ValueError, IndexError):
        print('❌ Invalid project selection')
        sys.exit(1)

    project_id = project['id']
    project_name = project['name']

    print('\n📚 Fetching test suites...')
    suites = normalize_list_response(client.get_suites(project_id), 'suites')

    suite_id = None
    suite_name = 'All Suites'
    if suites:
        print('\nAvailable test suites:')
        print('  0. All suites')
        for i, s in enumerate(suites, 1):
            print(f'  {i}. {s["name"]} (ID: {s["id"]})')

        try:
            s_choice = int(input('\nSelect suite number (0 for all): ').strip())
            if s_choice > 0:
                suite = suites[s_choice - 1]
                suite_id = suite['id']
                suite_name = suite['name']
        except (ValueError, IndexError):
            print('❌ Invalid suite selection')
            sys.exit(1)
    else:
        print('ℹ️ No suites returned (single-suite mode project maybe). Continue with all.')

    print('\n📚 Fetching sections...')
    sections = fetch_all_sections(client, project_id, suite_id=suite_id)
    by_id, children_map, section_path = build_section_maps(sections)

    selected_section_ids = None
    if sections:
        print('\nSection filter mode:')
        print('  0. All sections')
        print('  1. Quick select by section IDs')
        print('  2. Hierarchical step-by-step picker (recommended)')
        mode = input('Select mode (0/1/2): ').strip()

        if mode == '1':
            print('\nAvailable sections (ID -> Path):')
            for sid in sorted(section_path):
                print(f'  {sid} -> {section_path[sid]}')

            raw_ids = input('\nEnter section ID(s), comma/space separated: ').strip()
            try:
                base_ids = parse_section_selection(raw_ids)
            except ValueError:
                print('❌ Invalid section IDs format')
                sys.exit(1)

            missing = [sid for sid in base_ids if sid not in by_id]
            if missing:
                print(f'❌ Unknown section ID(s): {missing}')
                sys.exit(1)

            include_sub = input('Include subsections? (y/n): ').strip().lower() == 'y'
            selected_section_ids = collect_descendants(children_map, base_ids) if include_sub else set(base_ids)
            print(f'✅ Selected {len(selected_section_ids)} section(s)')
        elif mode == '2':
            selected_section_ids = choose_sections_interactive(by_id, children_map, section_path)
            if selected_section_ids is None:
                print('✅ Using all sections (no filtering)')
            else:
                print(f'✅ Selected {len(selected_section_ids)} section(s) via hierarchy')
        else:
            print('✅ Using all sections')
    else:
        print('ℹ️ No sections found, fallback to all cases')

    start = datetime.now()

    if selected_section_ids is None:
        print('\n📋 Fetching all cases from the suite...')
    else:
        print(f'\n📋 Fetching cases from {len(selected_section_ids)} selected sections...')

    all_cases = fetch_all_cases(
        client,
        project_id,
        suite_id=suite_id,
        selected_section_ids=selected_section_ids
    )

    case_rows = []
    auto_field_stats = {}
    for c in all_cases:
        sid = c.get('section_id')
        automation_id = extract_automation_id(c)
        has_automation = bool(automation_id)

        if AUTOMATION_ID_FIELD in c:
            auto_field_stats[AUTOMATION_ID_FIELD] = auto_field_stats.get(AUTOMATION_ID_FIELD, 0) + 1

        case_rows.append({
            'id': c['id'],
            'title': c.get('title', ''),
            'section_id': sid,
            'section': section_path.get(sid, f'Section {sid}'),
            'link': f'{testrail_url.rstrip("/")}/index.php?/cases/view/{c["id"]}',
            'automation_id': automation_id,
            'has_automation_id': has_automation,
        })

    print(f'📋 Candidate cases after section filter: {len(case_rows)}')
    print(f'🔍 Automation field presence stats: {auto_field_stats if auto_field_stats else "none"}')

    print('\n🏃 Fetching runs from plans...')
    runs_from_plans, run_to_plan = fetch_runs_from_plans(client, project_id, suite_id=suite_id)

    print('🏃 Fetching standalone runs...')
    standalone_runs = fetch_standalone_runs(client, project_id, suite_id=suite_id)

    merged = {**runs_from_plans, **standalone_runs}
    all_runs = list(merged.values())
    print(f'✅ Total unique runs considered: {len(all_runs)}')

    non_ci_runs, ci_runs = split_runs_ci_non_ci(all_runs, run_to_plan)

    raw_workers = input('\nTurbo workers count [default 4]: ').strip()
    try:
        workers = int(raw_workers) if raw_workers else 4
        workers = max(1, min(workers, 32))
    except ValueError:
        workers = 4

    cp_enabled = (input('Enable checkpoint CSV export? (Y/n): ').strip().lower() != 'n')
    if cp_enabled:
      raw_cp_every = input('Checkpoint every N completed runs [default 50]: ').strip()
      try:
          checkpoint_every_runs = int(raw_cp_every) if raw_cp_every else 50
          checkpoint_every_runs = max(1, checkpoint_every_runs)
    except ValueError:
        checkpoint_every_runs = 50
    else:
      checkpoint_every_runs = 0

    cache_enabled = (input('Enable in-memory get_tests cache? (Y/n): ').strip().lower() != 'n')

    checkpoint_base = f'testrail_v547_{datetime.now().strftime("%Y%m%d_%H%M%S")}'
    tests_cache = {}

    auto_cases = [c for c in case_rows if c['has_automation_id']]
    non_auto_cases = [c for c in case_rows if not c['has_automation_id']]
    print(f'🔧 Cases with automation_id: {len(auto_cases)}')
    print(f'📝 Cases without automation_id: {len(non_auto_cases)}')

    history_ids, in_open_ids = set(), set()

    if non_ci_runs:
        print('\n=== Stage 1/2: scanning NON-CI runs for all candidate cases ===')
        history_ids, in_open_ids = scan_runs_newest_first(
            client=client,
            candidate_cases=case_rows,
            all_runs=non_ci_runs,
            workers=workers,
            history_ids=history_ids,
            in_open_ids=in_open_ids,
            stage_label='stage1_non_ci',
            checkpoint_enabled=cp_enabled,
            checkpoint_every_runs=checkpoint_every_runs,
            checkpoint_base_name=checkpoint_base,
            tests_cache_enabled=cache_enabled,
            tests_cache=tests_cache
        )
    else:
        print('\nℹ️ Stage 1 skipped: no non-CI runs found.')

    auto_cases_stage2 = [c for c in auto_cases if c['id'] not in history_ids]
    print(f'🚀 Stage 2 candidate automated cases (unresolved after Stage 1): {len(auto_cases_stage2)}')

    if ci_runs and auto_cases_stage2:
        print('\n=== Stage 2/2: scanning CI runs ONLY for unresolved cases with automation_id ===')
        history_ids, in_open_ids = scan_runs_newest_first(
            client=client,
            candidate_cases=auto_cases_stage2,
            all_runs=ci_runs,
            workers=workers,
            history_ids=history_ids,
            in_open_ids=in_open_ids,
            stage_label='stage2_ci_unresolved_auto',
            checkpoint_enabled=cp_enabled,
            checkpoint_every_runs=checkpoint_every_runs,
            checkpoint_base_name=checkpoint_base,
            tests_cache_enabled=cache_enabled,
            tests_cache=tests_cache
        )
    else:
        if not ci_runs:
            print('\nℹ️ Stage 2 skipped: no CI runs found.')
        elif not auto_cases_stage2:
            print('\nℹ️ Stage 2 skipped: all automated cases were already resolved in Stage 1.')

    report_rows = build_final_report(case_rows, history_ids, in_open_ids)

    duration = (datetime.now() - start).total_seconds()
    total = len(report_rows)
    never_exec = sum(1 for r in report_rows if r['status'] == 'never_executed')
    has_hist = sum(1 for r in report_rows if r['status'] == 'has_history')
    in_open = sum(1 for r in report_rows if r['status'] == 'in_open_run_no_results')

    print('\n📈 Summary')
    print(f'   Project: {project_name}')
    print(f'   Suite: {suite_name}')
    print(f'   Candidate cases: {total}')
    print(f'   has_history: {has_hist}')
    print(f'   in_open_run_no_results: {in_open}')
    print(f'   never_executed: {never_exec}')
    print(f'   ⏱️ Time taken: {duration:.1f} sec')

    print('\nℹ️ Note:')
    print("   Cases without automation_id are not scanned in CI runs by design.")
    print("   'never_executed' means no results found in: NON-CI runs + CI runs(for unresolved automated cases).")

    scope = input('\nReport scope: (1) All checked cases, (2) Only never executed + safe to delete: ').strip()
    if scope == '2':
        export_rows = [
            r for r in report_rows
            if r['status'] == 'never_executed' and r['safe_to_delete_candidate'] == 'YES'
        ]
        print(f'📄 Exporting filtered rows: {len(export_rows)} / {len(report_rows)}')
    else:
        export_rows = report_rows
        print(f'📄 Exporting all rows: {len(export_rows)}')

    if not export_rows:
        print('⚠️ Nothing to export for selected scope.')
        return

    out_fmt = input('\nSave report as (1=CSV, 2=XLSX): ').strip()
    default_base = f'testrail_case_history_analysis_runoptimized_v547_{datetime.now().strftime("%Y%m%d_%H%M%S")}'
    custom_base = input(f'Enter report file name (without extension) [{default_base}]: ').strip()
    base = safe_filename_base(custom_base or default_base, default_base)

    if out_fmt == '2':
        save_to_xlsx(export_rows, f'{base}.xlsx')
    else:
        save_to_csv(export_rows, f'{base}.csv')


if __name__ == '__main__':
    try:
        main()
    except KeyboardInterrupt:
        print('\n⚠️ Interrupted by user.')
        sys.exit(130)
    except Exception as exc:
        print(f'\n❌ Fatal error: {exc}')
        cause = exc.__cause__
        while cause is not None:
            print(f'   ↳ caused by: {type(cause).__name__}: {cause}')
            cause = cause.__cause__
        sys.exit(1)
